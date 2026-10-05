#!/usr/bin/env python3
"""style_lint.py — verifies the human-writer skill profile (PROMPT/SKILLS/human-writer.md).

Profile enforced on every shipped text artefact in SUBMISSION/:
  * no em dash (U+2014)
  * no semicolon
  * no asterisk (markdown emphasis removed)
  * no hashtags
  * no banned vocabulary from the skill's list

Word matching is word-boundary based, so "discovery" does not trip "discover".
Reads markdown, csv and the text layer of shipped PPTX/PDF. Exits non-zero on any hit.
"""
import glob, io, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
SUB = os.path.join(os.path.dirname(HERE), "SUBMISSION")

BANNED = ["can", "may", "just", "that", "very", "really", "literally", "actually", "certainly",
          "probably", "basically", "could", "maybe", "delve", "embark", "enlightening",
          "esteemed", "shed light", "craft", "crafting", "imagine", "realm", "game-changer",
          "unlock", "discover", "skyrocket", "abyss", "not alone", "in a world where",
          "revolutionize", "disruptive", "utilize", "utilizing", "dive deep", "tapestry",
          "illuminate", "unveil", "pivotal", "intricate", "elucidate", "hence", "furthermore",
          "however", "harness", "exciting", "groundbreaking", "cutting-edge", "remarkable",
          "remains to be seen", "glimpse into", "navigating", "landscape", "stark", "testament",
          "in summary", "in conclusion", "moreover", "boost", "skyrocketing", "opened up",
          "powerful", "inquiries", "ever-evolving"]

PUNCT = [("\u2014", "em dash"), (";", "semicolon"), ("*", "asterisk")]


def text_of(path):
    """Extract comparable text. Binaries yield their text layer."""
    if path.endswith(".pptx"):
        from pptx import Presentation
        pr = Presentation(path)
        parts = []
        for sl in pr.slides:
            for sh in sl.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False):
                    for row in sh.table.rows:
                        for cell in row.cells:
                            parts.append(cell.text)
        return "\n".join(parts)
    if path.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(path).pages)
    return io.open(path, encoding="utf-8").read()


def scan(path):
    text = text_of(path)
    hits = []
    for ch, name in PUNCT:
        for m in re.finditer(re.escape(ch), text):
            line = text.count("\n", 0, m.start()) + 1
            ctx = text[max(0, m.start() - 45):m.start() + 45].replace("\n", " ")
            hits.append((name, line, ctx))
    for b in BANNED:
        pat = r"(?<![\w-])" + r"\s+".join(re.escape(w) for w in b.split()) + r"(?![\w-])"
        for m in re.finditer(pat, text, re.IGNORECASE):
            line = text.count("\n", 0, m.start()) + 1
            ctx = text[max(0, m.start() - 45):m.start() + 45].replace("\n", " ")
            hits.append((b, line, ctx))
    for m in re.finditer(r"#\w", text):
        line = text.count("\n", 0, m.start()) + 1
        hits.append(("hashtag", line, text[max(0, m.start() - 45):m.start() + 45]))
    return hits


def main():
    import json
    paths = []
    for pat in ("*.md", "*.csv", "*.txt", "*.pptx", "*.pdf"):
        paths += sorted(glob.glob(os.path.join(SUB, pat)))
    paths = [p for p in paths if not p.endswith(".pyc")]
    total = 0
    print("style profile: human-writer skill  |  scope: %s" % SUB)
    print("")
    for p in paths:
        hits = scan(p)
        total += len(hits)
        print("%-52s %s" % (os.path.basename(p), "CLEAN" if not hits else "%d HIT(S)" % len(hits)))
        for name, line, ctx in hits[:40]:
            print("    line %-5s %-16s %s" % (line, name, ctx.strip()[:100]))
    print("")
    print("TOTAL HITS: %d" % total)
    # machine-readable record for the manifest
    here = os.path.dirname(os.path.abspath(__file__))
    rec = {"profile": "human-writer skill (PROMPT/SKILLS/human-writer.md)",
           "scope": "SUBMISSION/",
           "rules": ["no em dash", "no semicolon", "no asterisk", "no hashtag", "no banned vocabulary"],
           "files": {}, "total_hits": total, "pass": total == 0}
    for p in paths:
        rec["files"][os.path.basename(p)] = len(scan(p))
    io.open(os.path.join(os.path.dirname(here), "state", "style_lint.json"), "w",
            encoding="utf-8").write(json.dumps(rec, indent=2) + "\n")
    return 1 if total else 0


if __name__ == "__main__":
    sys.exit(main())
